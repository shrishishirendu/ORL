"""Builds the ``.xlsx`` workbook ``GET /admin/data/template`` returns.

Column names/order are imported from ``app.services.admin_data.upload``
(the parser's own source-of-truth constants) rather than duplicated here,
so the template and the parser can never quietly drift apart.
"""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.services.admin_data.upload import (
    SHIFTS_COLUMNS,
    SHIFTS_SHEET,
    WORKERS_COLUMNS,
    WORKERS_SHEET,
)

# One clearly-fake example row per sheet -- a placeholder to show the
# expected shape, not real data (see this module's docstring / the task
# brief's "clearly a placeholder example, not real data").
_WORKERS_EXAMPLE_ROW = ["Jane Example", "nursing, first_aid", "north", "NTH-01", "true", "EMP-0001"]
_SHIFTS_EXAMPLE_ROW = ["2026-04-01", "08:00", "16:00", "nursing", "NTH-01", "false"]

_LEGEND_HEADER = ["Sheet", "Column", "Required?", "Notes"]
_LEGEND_ROWS = [
    ("Workers", "name", "required", "Full name."),
    ("Workers", "skills", "required", "Comma-separated, e.g. 'nursing, first_aid'."),
    (
        "Workers",
        "region",
        "optional",
        "Free-text region; Tier 1 eligibility normally resolves region from "
        "home_site_code's Site.region instead, so this is mostly an override/fallback.",
    ),
    (
        "Workers",
        "home_site_code",
        "required",
        "Must match an existing Site.code -- see GET /sites.",
    ),
    ("Workers", "active", "optional (default true)", "true/false/1/0/yes/no."),
    (
        "Workers",
        "employee_code",
        "optional",
        "Stable natural key used to match this row to an existing worker on a future "
        "upload. If omitted, matching falls back to an exact case-insensitive name "
        "match instead, which is less reliable (flagged as such in the upload result).",
    ),
    ("Shifts", "date", "required", "ISO date, e.g. 2026-04-01."),
    ("Shifts", "start_time", "required", "HH:MM, 24-hour."),
    (
        "Shifts",
        "end_time",
        "required",
        "HH:MM, 24-hour. If <= start_time, treated as an overnight shift ending the next day.",
    ),
    ("Shifts", "required_skill", "required", "Must match a skill string used in Worker.skills."),
    ("Shifts", "site_code", "required", "Must match an existing Site.code -- see GET /sites."),
    (
        "Shifts",
        "is_multi_stop",
        "must be false or blank",
        "Multi-stop shifts need Job rows, which this upload doesn't manage -- leave "
        "false/blank, or use the existing seed/DB path for multi-stop shifts.",
    ),
]


def _autosize_columns(ws: Worksheet) -> None:
    for col_cells in ws.columns:
        lengths = [len(str(c.value)) for c in col_cells if c.value is not None]
        width = min(max((max(lengths, default=10)) + 2, 10), 60)
        ws.column_dimensions[col_cells[0].column_letter].width = width


def build_template_workbook() -> bytes:
    """Build the template workbook and return its raw ``.xlsx`` bytes."""
    wb = Workbook()

    ws_workers = wb.active
    ws_workers.title = WORKERS_SHEET
    ws_workers.append(WORKERS_COLUMNS)
    ws_workers.append(_WORKERS_EXAMPLE_ROW)

    ws_shifts = wb.create_sheet(SHIFTS_SHEET)
    ws_shifts.append(SHIFTS_COLUMNS)
    ws_shifts.append(_SHIFTS_EXAMPLE_ROW)

    ws_legend = wb.create_sheet("Legend")
    ws_legend.append(_LEGEND_HEADER)
    for row in _LEGEND_ROWS:
        ws_legend.append(row)

    for ws in (ws_workers, ws_shifts, ws_legend):
        _autosize_columns(ws)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
